import io
from datetime import datetime
from decimal import Decimal
from django.http import HttpResponse
from django.utils import timezone

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors, pagesizes
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer, KeepTogether, PageBreak
)
from reportlab.pdfgen import canvas

from apps.financial.services.financial_service import FinancialReportService
from apps.financial.services.reports_service import ReportsService


class NumberedCanvas(canvas.Canvas):
    """
    Two-pass canvas to dynamically compute and print total page numbers: 'Page X of Y'.
    """
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            super().showPage()
        super().save()

    def draw_page_number(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#78716C"))
        
        # Draw bottom footer
        page_str = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(self._pagesize[0] - 36, 20, page_str)
        self.drawString(36, 20, "EasyServe Restaurant Management SaaS — Confidential & Proprietary Report")
        self.setStrokeColor(colors.HexColor("#E7E5E4"))
        self.setLineWidth(0.5)
        self.line(36, 32, self._pagesize[0] - 36, 32)
        self.restoreState()


class ExportService:
    """
    Generates professional, printable, audit-ready Excel and PDF reports.
    """

    EMERALD = "063B2E"
    GOLD = "D4A72C"
    OFF_WHITE = "F7F7F4"
    BORDER_GRAY = "D6D3D1"

    @classmethod
    def get_report_dataset(cls, report_type, restaurant, start_date, end_date):
        t = (report_type or "pnl").lower()
        if t == "pnl":
            return ReportsService.get_pnl_report(restaurant, start_date, end_date)
        elif t == "sales":
            return ReportsService.get_sales_report(restaurant, start_date, end_date)
        elif t == "cogs":
            return ReportsService.get_cogs_report(restaurant, start_date, end_date)
        elif t == "products":
            return ReportsService.get_product_performance_report(restaurant, start_date, end_date)
        elif t == "expenses":
            return ReportsService.get_expense_report(restaurant, start_date, end_date)
        elif t == "wastage":
            return ReportsService.get_wastage_report(restaurant, start_date, end_date)
        elif t == "inventory":
            return ReportsService.get_inventory_report(restaurant)
        elif t == "movements":
            return ReportsService.get_stock_movement_report(restaurant, start_date, end_date)
        elif t == "purchases":
            return ReportsService.get_purchase_report(restaurant, start_date, end_date)
        elif t == "reconciliation":
            return ReportsService.get_reconciliation_report(restaurant, start_date, end_date)
        else:
            return ReportsService.get_pnl_report(restaurant, start_date, end_date)

    # =========================================================================
    # EXCEL EXPORT ENGINE
    # =========================================================================

    @classmethod
    def generate_excel(cls, report_type, restaurant, start_date, end_date):
        data = cls.get_report_dataset(report_type, restaurant, start_date, end_date)
        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        t = (report_type or "pnl").lower()
        rest_name = restaurant.name if restaurant else "EasyServe Restaurant"
        period_str = f"{start_date.strftime('%d %b %Y')} - {end_date.strftime('%d %b %Y')}" if start_date and end_date else "Current Date"
        gen_time = FinancialReportService.get_local_now().strftime("%Y-%m-%d %H:%M:%S")

        if t == "pnl":
            cls._build_excel_pnl(wb, data, rest_name, period_str, gen_time)
        elif t == "sales":
            cls._build_excel_sales(wb, data, rest_name, period_str, gen_time)
        elif t == "cogs":
            cls._build_excel_cogs(wb, data, rest_name, period_str, gen_time)
        elif t == "products":
            cls._build_excel_products(wb, data, rest_name, period_str, gen_time)
        elif t == "expenses":
            cls._build_excel_expenses(wb, data, rest_name, period_str, gen_time)
        elif t == "wastage":
            cls._build_excel_wastage(wb, data, rest_name, period_str, gen_time)
        elif t == "inventory":
            cls._build_excel_inventory(wb, data, rest_name, period_str, gen_time)
        elif t == "movements":
            cls._build_excel_movements(wb, data, rest_name, period_str, gen_time)
        elif t == "purchases":
            cls._build_excel_purchases(wb, data, rest_name, period_str, gen_time)
        elif t == "reconciliation":
            cls._build_excel_reconciliation(wb, data, rest_name, period_str, gen_time)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"EasyServe_{t.upper()}_{start_date.isoformat()}_to_{end_date.isoformat()}.xlsx"
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response

    @classmethod
    def _apply_excel_header(cls, ws, title, rest_name, period_str, gen_time, max_col=6):
        fill_dark = PatternFill(start_color=cls.EMERALD, end_color=cls.EMERALD, fill_type="solid")
        fill_sub = PatternFill(start_color="0A4F3E", end_color="0A4F3E", fill_type="solid")
        font_title = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
        font_sub = Font(name="Calibri", size=11, bold=False, color="D4A72C")
        font_meta = Font(name="Calibri", size=9, italic=True, color="D6D3D1")

        # Row 1: Restaurant & Main Title
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
        cell1 = ws.cell(row=1, column=1, value=f"{rest_name.upper()} — {title}")
        cell1.font = font_title
        cell1.fill = fill_dark
        cell1.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[1].height = 36

        # Row 2: Period & Generation metadata
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col)
        cell2 = ws.cell(row=2, column=1, value=f"Reporting Period: {period_str}   |   Generated: {gen_time} (Asia/Karachi)")
        cell2.font = font_sub
        cell2.fill = fill_sub
        cell2.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[2].height = 22

        ws.row_dimensions[3].height = 10  # blank spacer

    @classmethod
    def _style_excel_table(cls, ws, start_row, header_cols, data_rows, totals_row=None):
        header_fill = PatternFill(start_color=cls.EMERALD, end_color=cls.EMERALD, fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        zebra_fill = PatternFill(start_color="F9F9F6", end_color="F9F9F6", fill_type="solid")
        thin_border = Border(
            left=Side(style='thin', color="E5E7EB"),
            right=Side(style='thin', color="E5E7EB"),
            top=Side(style='thin', color="E5E7EB"),
            bottom=Side(style='thin', color="E5E7EB"),
        )
        double_bottom = Border(
            top=Side(style='thin', color="063B2E"),
            bottom=Side(style='double', color="063B2E"),
        )

        # Header
        ws.row_dimensions[start_row].height = 24
        for col_idx, h_text in enumerate(header_cols, start=1):
            cell = ws.cell(row=start_row, column=col_idx, value=h_text)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        # Data Rows
        current_row = start_row + 1
        for r_idx, row_values in enumerate(data_rows):
            ws.row_dimensions[current_row].height = 20
            use_zebra = (r_idx % 2 == 1)
            for c_idx, val in enumerate(row_values, start=1):
                cell = ws.cell(row=current_row, column=c_idx, value=val)
                cell.border = thin_border
                cell.font = Font(name="Calibri", size=10)
                if use_zebra:
                    cell.fill = zebra_fill

                # Formatting
                if isinstance(val, (int, float, Decimal)):
                    if "Rs." in str(header_cols[c_idx - 1]) or "Amount" in header_cols[c_idx - 1] or "Sales" in header_cols[c_idx - 1] or "Cost" in header_cols[c_idx - 1] or "Profit" in header_cols[c_idx - 1] or "Spend" in header_cols[c_idx - 1] or "Value" in header_cols[c_idx - 1]:
                        cell.number_format = '"Rs. "#,##0.00'
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    elif "%" in header_cols[c_idx - 1] or "Margin" in header_cols[c_idx - 1] or "Percentage" in header_cols[c_idx - 1]:
                        cell.number_format = '0.0"%"'
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    else:
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                elif isinstance(val, str) and (val.startswith("202") or val.startswith("ORD") or val.startswith("EXP") or val.startswith("PO")):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
            current_row += 1

        # Totals Row
        if totals_row:
            ws.row_dimensions[current_row].height = 22
            for c_idx, val in enumerate(totals_row, start=1):
                cell = ws.cell(row=current_row, column=c_idx, value=val)
                cell.font = Font(name="Calibri", size=10, bold=True, color="063B2E")
                cell.border = double_bottom
                cell.fill = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")
                if isinstance(val, (int, float, Decimal)):
                    cell.number_format = '"Rs. "#,##0.00'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
            current_row += 1

        # Auto column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # Freeze Header
        ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
        # Enable Auto-filters
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(header_cols))}{current_row - 1}"

    @classmethod
    def _build_excel_pnl(cls, wb, data, rest_name, period_str, gen_time):
        ws = wb.create_sheet(title="P&L Statement")
        cls._apply_excel_header(ws, "Profit & Loss Statement", rest_name, period_str, gen_time, max_col=4)

        ov = data.get("overview", {})
        rev = ov.get("revenue", {})
        costs = ov.get("costs", {})
        prof = ov.get("profitability", {})

        pnl_lines = [
            ["1. REVENUE", "", "", ""],
            ["   Gross Sales", float(rev.get("gross_sales", 0)), "", "Total realized sales"],
            ["   Discounts & Deductions", float(rev.get("discounts", 0)), "", "Standard discounts"],
            ["   Net Sales", float(rev.get("net_sales", 0)), "100.0%", "Base operational revenue"],
            ["", "", "", ""],
            ["2. COST OF GOODS SOLD (COGS)", "", "", ""],
            ["   Cost of Ingredients (Historical)", float(costs.get("cogs", 0)), f"{float(costs.get('food_cost_percentage', 0)):.1f}%", "Snapshotted ingredient consumption"],
            ["", "", "", ""],
            ["3. GROSS PROFIT", float(prof.get("gross_profit", 0)), f"{float(prof.get('gross_margin', 0)):.1f}%", "Net Sales minus COGS"],
            ["", "", "", ""],
            ["4. OPERATING OVERHEADS & LOSSES", "", "", ""],
            ["   Operating Expenses (Overhead)", float(costs.get("operating_expenses", 0)), "", f"{costs.get('expense_count', 0)} active expense logs"],
            ["   Stock Wastage (Losses)", float(costs.get("wastage", 0)), "", f"{costs.get('wastage_count', 0)} logged incidents"],
            ["", "", "", ""],
            ["5. NET OPERATING RESULT", float(prof.get("operating_result", 0)), f"{float(prof.get('net_margin', 0)):.1f}%", "Final bottom line profit/loss"],
        ]

        cls._style_excel_table(
            ws,
            start_row=4,
            header_cols=["Financial Statement Line Item", "Amount (PKR)", "% of Sales", "Audit Notes"],
            data_rows=pnl_lines
        )

    @classmethod
    def _build_excel_sales(cls, wb, data, rest_name, period_str, gen_time):
        ws = wb.create_sheet(title="Sales Orders")
        cls._apply_excel_header(ws, "Sales Activity Ledger", rest_name, period_str, gen_time, max_col=10)

        rows = [
            [
                o["order_number"], o["date"], o["time"], o["table_number"],
                o["order_type"], o["status"], o["payment_mode"],
                o["net_amount"], o["cogs"], o["gross_profit"]
            ]
            for o in data.get("orders", [])
        ]
        sum_data = data.get("summary", {})
        totals = ["TOTALS", "", "", "", "", "", f"{sum_data.get('total_orders', 0)} Orders", sum_data.get("total_net_sales", 0), sum_data.get("total_cogs", 0), sum_data.get("total_gross_profit", 0)]

        cls._style_excel_table(
            ws,
            start_row=4,
            header_cols=["Order #", "Date", "Time", "Table", "Type", "Status", "Payment", "Net Sales", "COGS", "Gross Profit"],
            data_rows=rows,
            totals_row=totals
        )

    @classmethod
    def _build_excel_cogs(cls, wb, data, rest_name, period_str, gen_time):
        ws = wb.create_sheet(title="COGS Analysis")
        cls._apply_excel_header(ws, "Food Cost & COGS Breakdown", rest_name, period_str, gen_time, max_col=7)

        rows = [
            [
                p.get("name"), p.get("category_name", "General"),
                p.get("units_sold", 0), float(p.get("revenue", 0)),
                float(p.get("cogs", 0)), float(p.get("gross_profit", 0)),
                float(p.get("food_cost_percentage", 0))
            ]
            for p in data.get("items", [])
        ]
        sum_data = data.get("summary", {})
        totals = ["TOTALS", "", f"{sum_data.get('total_menu_items_sold', 0)} Items", sum_data.get("total_revenue", 0), sum_data.get("total_cogs", 0), float(sum_data.get("total_revenue", 0)) - float(sum_data.get("total_cogs", 0)), sum_data.get("overall_food_cost_pct", 0)]

        cls._style_excel_table(
            ws,
            start_row=4,
            header_cols=["Menu Item", "Category", "Units Sold", "Revenue", "Historical COGS", "Gross Profit", "Food Cost %"],
            data_rows=rows,
            totals_row=totals
        )

    @classmethod
    def _build_excel_products(cls, wb, data, rest_name, period_str, gen_time):
        ws = wb.create_sheet(title="Product Performance")
        cls._apply_excel_header(ws, "Menu Item Performance", rest_name, period_str, gen_time, max_col=7)

        rows = [
            [
                p.get("name"), p.get("category_name", "General"),
                p.get("units_sold", 0), float(p.get("revenue", 0)),
                float(p.get("cogs", 0)), float(p.get("gross_profit", 0)),
                float(p.get("food_cost_percentage", 0))
            ]
            for p in data.get("products", [])
        ]

        cls._style_excel_table(
            ws,
            start_row=4,
            header_cols=["Menu Item", "Category", "Units Sold", "Revenue", "COGS", "Gross Profit", "Food Cost %"],
            data_rows=rows
        )

    @classmethod
    def _build_excel_expenses(cls, wb, data, rest_name, period_str, gen_time):
        ws = wb.create_sheet(title="Operating Expenses")
        cls._apply_excel_header(ws, "Operating Overhead Expenses", rest_name, period_str, gen_time, max_col=7)

        rows = [
            [
                e["expense_number"], e["expense_date"], e["category_name"],
                e["title"], e["vendor_payee"], e["payment_method"], e["amount"]
            ]
            for e in data.get("expenses", [])
        ]
        sum_data = data.get("summary", {})
        totals = ["TOTALS", "", "", f"{sum_data.get('expense_count', 0)} Expenses", "", "", sum_data.get("total_expenses", 0)]

        cls._style_excel_table(
            ws,
            start_row=4,
            header_cols=["Expense #", "Date", "Category", "Title", "Vendor/Payee", "Payment Method", "Amount (PKR)"],
            data_rows=rows,
            totals_row=totals
        )

    @classmethod
    def _build_excel_wastage(cls, wb, data, rest_name, period_str, gen_time):
        ws = wb.create_sheet(title="Stock Wastage")
        cls._apply_excel_header(ws, "Wastage & Spoilage Log", rest_name, period_str, gen_time, max_col=8)

        rows = [
            [
                w["date"], w["item_name"], w["quantity"], w["uom"],
                w["reason"], w["unit_cost"], w["total_cost"], w["notes"]
            ]
            for w in data.get("wastage_logs", [])
        ]
        sum_data = data.get("summary", {})
        totals = ["TOTALS", "", "", "", f"{sum_data.get('incident_count', 0)} Incidents", "", sum_data.get("total_loss", 0), ""]

        cls._style_excel_table(
            ws,
            start_row=4,
            header_cols=["Date", "Inventory Item", "Quantity", "UOM", "Reason", "Unit Cost", "Total Loss", "Notes"],
            data_rows=rows,
            totals_row=totals
        )

    @classmethod
    def _build_excel_inventory(cls, wb, data, rest_name, period_str, gen_time):
        ws = wb.create_sheet(title="Inventory Valuation")
        cls._apply_excel_header(ws, "Stock Valuation & Levels", rest_name, period_str, gen_time, max_col=8)

        rows = [
            [
                i["name"], i["sku"], i["category_name"],
                i["current_stock"], i["uom"], i["cost_per_unit"],
                i["stock_value"], i["status"].replace("_", " ")
            ]
            for i in data.get("items", [])
        ]
        sum_data = data.get("summary", {})
        totals = ["TOTALS", "", f"{sum_data.get('total_items', 0)} Items", "", "", "", sum_data.get("total_stock_value", 0), ""]

        cls._style_excel_table(
            ws,
            start_row=4,
            header_cols=["Item Name", "SKU", "Category", "Current Stock", "UOM", "Cost/Unit", "Stock Value", "Status"],
            data_rows=rows,
            totals_row=totals
        )

    @classmethod
    def _build_excel_movements(cls, wb, data, rest_name, period_str, gen_time):
        ws = wb.create_sheet(title="Stock Ledger")
        cls._apply_excel_header(ws, "Stock Movement Audit Ledger", rest_name, period_str, gen_time, max_col=8)

        rows = [
            [
                m["date"], m["time"], m["item_name"], m["movement_type"],
                m["quantity_delta"], m["balance_after"], m["uom"], m["reference_note"]
            ]
            for m in data.get("movements", [])
        ]

        cls._style_excel_table(
            ws,
            start_row=4,
            header_cols=["Date", "Time", "Item Name", "Movement Type", "Quantity Delta", "Balance After", "UOM", "Reference Note"],
            data_rows=rows
        )

    @classmethod
    def _build_excel_purchases(cls, wb, data, rest_name, period_str, gen_time):
        ws = wb.create_sheet(title="Purchases")
        cls._apply_excel_header(ws, "Purchase Orders & Stock In", rest_name, period_str, gen_time, max_col=7)

        rows = [
            [
                p["purchase_number"], p["invoice_number"], p["supplier_name"],
                p["purchase_date"], p["status"], p["subtotal"], p["total_amount"]
            ]
            for p in data.get("purchases", [])
        ]
        sum_data = data.get("summary", {})
        totals = ["TOTALS", "", "", f"{sum_data.get('total_purchase_count', 0)} POs", "", "", sum_data.get("total_spend", 0)]

        cls._style_excel_table(
            ws,
            start_row=4,
            header_cols=["PO #", "Invoice #", "Supplier", "Purchase Date", "Status", "Subtotal", "Total Amount"],
            data_rows=rows,
            totals_row=totals
        )

    @classmethod
    def _build_excel_reconciliation(cls, wb, data, rest_name, period_str, gen_time):
        ws = wb.create_sheet(title="Audit Reconciliation")
        cls._apply_excel_header(ws, "Audit & Cash Reconciliation", rest_name, period_str, gen_time, max_col=4)

        rec = data.get("reconciliation", {})
        rev_rec = rec.get("revenue_reconciliation", {})
        cogs_rec = rec.get("cogs_reconciliation", {})
        exp_rec = rec.get("expense_reconciliation", {})

        rows = [
            ["1. REVENUE AUDIT RECONCILIATION", "", "", ""],
            ["   Orders Gross Sales Total", float(rev_rec.get("orders_gross_revenue", 0)), "", "Source: Orders ledger"],
            ["   Confirmed Payments Total", float(rev_rec.get("payments_total", 0)), "", "Source: Payment transactions"],
            ["   Revenue Variance / Difference", float(rev_rec.get("difference", 0)), "", f"Status: {rev_rec.get('status')}"],
            ["", "", "", ""],
            ["2. COGS AUDIT RECONCILIATION", "", "", ""],
            ["   Order-Level Snapshotted COGS Total", float(cogs_rec.get("order_cogs_total", 0)), "", "Source: Orders.total_cogs"],
            ["   Item-Level Snapshotted COGS Total", float(cogs_rec.get("item_cogs_total", 0)), "", "Source: OrderItem.unit_cost_at_order"],
            ["   COGS Variance / Difference", float(cogs_rec.get("difference", 0)), "", f"Status: {cogs_rec.get('status')}"],
            ["", "", "", ""],
            ["3. EXPENSES AUDIT RECONCILIATION", "", "", ""],
            ["   Active Expense Records Sum", float(exp_rec.get("expense_records_sum", 0)), "", "Source: Active Expense table"],
            ["   P&L Expense Figure", float(exp_rec.get("expense_summary_total", 0)), "", "Source: P&L Statement"],
            ["   Expense Variance / Difference", float(exp_rec.get("difference", 0)), "", f"Status: {exp_rec.get('status')}"],
        ]

        cls._style_excel_table(
            ws,
            start_row=4,
            header_cols=["Audit Reconciliation Category", "Amount (PKR)", "Status", "Audit Traceability"],
            data_rows=rows
        )

    # =========================================================================
    # PDF EXPORT ENGINE (REPORTLAB)
    # =========================================================================

    @classmethod
    def generate_pdf(cls, report_type, restaurant, start_date, end_date):
        data = cls.get_report_dataset(report_type, restaurant, start_date, end_date)
        t = (report_type or "pnl").lower()
        rest_name = restaurant.name if restaurant else "EasyServe Restaurant"
        period_str = f"{start_date.strftime('%d %b %Y')} - {end_date.strftime('%d %b %Y')}" if start_date and end_date else "Current Date"
        gen_time = FinancialReportService.get_local_now().strftime("%Y-%m-%d %H:%M:%S")

        buffer = io.BytesIO()
        # Wide reports -> Landscape, Summary reports -> Portrait
        is_landscape = t in ["sales", "movements", "inventory", "purchases", "cogs", "products"]
        page_size = pagesizes.landscape(pagesizes.A4) if is_landscape else pagesizes.portrait(pagesizes.A4)

        doc = SimpleDocTemplate(
            buffer,
            pagesize=page_size,
            leftMargin=36,
            rightMargin=36,
            topMargin=36,
            bottomMargin=45
        )

        styles = getSampleStyleSheet()
        normal_style = styles["Normal"]
        
        title_style = ParagraphStyle(
            "ReportTitle",
            parent=normal_style,
            fontName="Helvetica-Bold",
            fontSize=16,
            textColor=colors.HexColor(f"#{cls.EMERALD}"),
            leading=20,
        )
        subtitle_style = ParagraphStyle(
            "ReportSubtitle",
            parent=normal_style,
            fontName="Helvetica",
            fontSize=9,
            textColor=colors.HexColor("#78716C"),
            leading=12,
        )

        elements = []

        # 1. Header Banner Box
        header_data = [
            [
                Paragraph(f"<b>{rest_name}</b><br/><font color='#D4A72C' size='12'><b>{data.get('title', 'Management Report')}</b></font>", title_style),
                Paragraph(f"<b>Period:</b> {period_str}<br/><b>Generated:</b> {gen_time}<br/><b>System:</b> EasyServe SaaS", subtitle_style)
            ]
        ]
        header_table = Table(header_data, colWidths=[doc.width * 0.6, doc.width * 0.4])
        header_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F7F7F4")),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor(f"#{cls.EMERALD}")),
            ('PADDING', (0, 0), (-1, -1), 8),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))
        elements.append(header_table)
        elements.append(Spacer(1, 14))

        # 2. Body Tables for Report Types
        if t == "pnl":
            cls._build_pdf_pnl(elements, data, doc.width, styles)
        elif t == "sales":
            cls._build_pdf_sales(elements, data, doc.width, styles)
        elif t == "cogs":
            cls._build_pdf_cogs(elements, data, doc.width, styles)
        elif t == "products":
            cls._build_pdf_products(elements, data, doc.width, styles)
        elif t == "expenses":
            cls._build_pdf_expenses(elements, data, doc.width, styles)
        elif t == "wastage":
            cls._build_pdf_wastage(elements, data, doc.width, styles)
        elif t == "inventory":
            cls._build_pdf_inventory(elements, data, doc.width, styles)
        elif t == "movements":
            cls._build_pdf_movements(elements, data, doc.width, styles)
        elif t == "purchases":
            cls._build_pdf_purchases(elements, data, doc.width, styles)
        elif t == "reconciliation":
            cls._build_pdf_reconciliation(elements, data, doc.width, styles)

        doc.build(elements, canvasmaker=NumberedCanvas)
        buffer.seek(0)

        filename = f"EasyServe_{t.upper()}_{start_date.isoformat()}_to_{end_date.isoformat()}.pdf"
        response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response

    @classmethod
    def _build_pdf_pnl(cls, elements, data, width, styles):
        ov = data.get("overview", {})
        rev = ov.get("revenue", {})
        costs = ov.get("costs", {})
        prof = ov.get("profitability", {})

        table_data = [
            ["Financial Statement Line Item", "Amount (PKR)", "% Sales", "Notes"],
            ["Gross Sales", f"Rs. {float(rev.get('gross_sales', 0)):,.2f}", "", f"{rev.get('orders_count', 0)} orders"],
            ["Discounts & Deductions", f"Rs. {float(rev.get('discounts', 0)):,.2f}", "", "-"],
            ["Net Sales", f"Rs. {float(rev.get('net_sales', 0)):,.2f}", "100.0%", "Base Sales"],
            ["Cost of Goods Sold (Historical COGS)", f"Rs. {float(costs.get('cogs', 0)):,.2f}", f"{float(costs.get('food_cost_percentage', 0)):.1f}%", "Ingredients consumed"],
            ["Gross Profit", f"Rs. {float(prof.get('gross_profit', 0)):,.2f}", f"{float(prof.get('gross_margin', 0)):.1f}%", "Net Sales minus COGS"],
            ["Operating Expenses (Overhead)", f"Rs. {float(costs.get('operating_expenses', 0)):,.2f}", "", f"{costs.get('expense_count', 0)} logs"],
            ["Stock Wastage Loss", f"Rs. {float(costs.get('wastage', 0)):,.2f}", "", f"{costs.get('wastage_count', 0)} incidents"],
            ["Net Operating Result", f"Rs. {float(prof.get('operating_result', 0)):,.2f}", f"{float(prof.get('net_margin', 0)):.1f}%", "Final bottom line profit"],
        ]

        t = Table(table_data, colWidths=[width * 0.45, width * 0.22, width * 0.13, width * 0.20])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(f"#{cls.EMERALD}")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('ALIGN', (1, 0), (2, -1), 'RIGHT'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9F9F6")]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
            ('FONTNAME', (0, 3), (-1, 3), 'Helvetica-Bold'),  # Net sales bold
            ('FONTNAME', (0, 5), (-1, 5), 'Helvetica-Bold'),  # Gross profit bold
            ('FONTNAME', (0, 8), (-1, 8), 'Helvetica-Bold'),  # Net profit bold
            ('BACKGROUND', (0, 8), (-1, 8), colors.HexColor("#ECFDF5")),
            ('TEXTCOLOR', (0, 8), (-1, 8), colors.HexColor("#063B2E")),
            ('PADDING', (0, 0), (-1, -1), 6),
        ]))
        elements.append(t)

    @classmethod
    def _build_pdf_sales(cls, elements, data, width, styles):
        orders = data.get("orders", [])[:100]  # First 100 on PDF
        table_data = [["Order #", "Date", "Time", "Type", "Status", "Payment", "Sales (PKR)", "COGS (PKR)", "Gross Profit"]]
        for o in orders:
            table_data.append([
                o["order_number"], o["date"], o["time"], o["order_type"],
                o["status"], o["payment_mode"],
                f"Rs. {o['net_amount']:,.2f}", f"Rs. {o['cogs']:,.2f}", f"Rs. {o['gross_profit']:,.2f}"
            ])

        sum_data = data.get("summary", {})
        table_data.append([
            "TOTALS", "", "", "", "", f"{sum_data.get('total_orders', 0)} Orders",
            f"Rs. {sum_data.get('total_net_sales', 0):,.2f}",
            f"Rs. {sum_data.get('total_cogs', 0):,.2f}",
            f"Rs. {sum_data.get('total_gross_profit', 0):,.2f}"
        ])

        col_w = [width * 0.11, width * 0.10, width * 0.08, width * 0.09, width * 0.10, width * 0.10, width * 0.14, width * 0.14, width * 0.14]
        t = Table(table_data, colWidths=col_w, repeatRows=1)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(f"#{cls.EMERALD}")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (6, 0), (-1, -1), 'RIGHT'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor("#F9F9F6")]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
            ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
            ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor("#F0FDF4")),
            ('PADDING', (0, 0), (-1, -1), 4),
        ]))
        elements.append(t)

    @classmethod
    def _build_pdf_cogs(cls, elements, data, width, styles):
        items = data.get("items", [])[:100]
        table_data = [["Menu Item", "Category", "Sold", "Revenue (PKR)", "Historical COGS", "Gross Profit", "Food Cost %"]]
        for p in items:
            table_data.append([
                p.get("name"), p.get("category_name", "General"), str(p.get("units_sold", 0)),
                f"Rs. {float(p.get('revenue', 0)):,.2f}", f"Rs. {float(p.get('cogs', 0)):,.2f}",
                f"Rs. {float(p.get('gross_profit', 0)):,.2f}", f"{float(p.get('food_cost_percentage', 0)):.1f}%"
            ])

        col_w = [width * 0.25, width * 0.15, width * 0.08, width * 0.14, width * 0.14, width * 0.14, width * 0.10]
        t = Table(table_data, colWidths=col_w, repeatRows=1)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(f"#{cls.EMERALD}")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (2, 0), (-1, -1), 'RIGHT'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9F9F6")]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
            ('PADDING', (0, 0), (-1, -1), 4),
        ]))
        elements.append(t)

    @classmethod
    def _build_pdf_products(cls, elements, data, width, styles):
        cls._build_pdf_cogs(elements, {"items": data.get("products", [])}, width, styles)

    @classmethod
    def _build_pdf_expenses(cls, elements, data, width, styles):
        exp = data.get("expenses", [])
        table_data = [["Expense #", "Date", "Category", "Title", "Vendor / Payee", "Payment Method", "Amount (PKR)"]]
        for e in exp:
            table_data.append([
                e["expense_number"], e["expense_date"], e["category_name"],
                e["title"], e["vendor_payee"], e["payment_method"],
                f"Rs. {e['amount']:,.2f}"
            ])

        sum_data = data.get("summary", {})
        table_data.append(["TOTALS", "", "", f"{sum_data.get('expense_count', 0)} Expenses", "", "", f"Rs. {sum_data.get('total_expenses', 0):,.2f}"])

        col_w = [width * 0.14, width * 0.11, width * 0.14, width * 0.22, width * 0.15, width * 0.11, width * 0.13]
        t = Table(table_data, colWidths=col_w, repeatRows=1)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(f"#{cls.EMERALD}")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (6, 0), (6, -1), 'RIGHT'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor("#F9F9F6")]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
            ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
            ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor("#F0FDF4")),
            ('PADDING', (0, 0), (-1, -1), 4),
        ]))
        elements.append(t)

    @classmethod
    def _build_pdf_wastage(cls, elements, data, width, styles):
        wst = data.get("wastage_logs", [])
        table_data = [["Date", "Inventory Item", "Quantity", "UOM", "Reason", "Unit Cost", "Total Loss (PKR)"]]
        for w in wst:
            table_data.append([
                w["date"], w["item_name"], f"{w['quantity']:.2f}", w["uom"],
                w["reason"], f"Rs. {w['unit_cost']:,.2f}", f"Rs. {w['total_cost']:,.2f}"
            ])

        sum_data = data.get("summary", {})
        table_data.append(["TOTALS", "", "", "", f"{sum_data.get('incident_count', 0)} Incidents", "", f"Rs. {sum_data.get('total_loss', 0):,.2f}"])

        col_w = [width * 0.12, width * 0.24, width * 0.10, width * 0.08, width * 0.18, width * 0.13, width * 0.15]
        t = Table(table_data, colWidths=col_w, repeatRows=1)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(f"#{cls.EMERALD}")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (2, 0), (2, -1), 'RIGHT'),
            ('ALIGN', (5, 0), (6, -1), 'RIGHT'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor("#F9F9F6")]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
            ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
            ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor("#FEF2F2")),
            ('PADDING', (0, 0), (-1, -1), 4),
        ]))
        elements.append(t)

    @classmethod
    def _build_pdf_inventory(cls, elements, data, width, styles):
        items = data.get("items", [])
        table_data = [["Item Name", "SKU", "Category", "Current Stock", "UOM", "Cost / Unit", "Stock Value", "Status"]]
        for i in items:
            table_data.append([
                i["name"], i["sku"], i["category_name"],
                f"{i['current_stock']:.2f}", i["uom"],
                f"Rs. {i['cost_per_unit']:,.2f}", f"Rs. {i['stock_value']:,.2f}",
                i["status"].replace("_", " ")
            ])

        sum_data = data.get("summary", {})
        table_data.append(["TOTALS", "", f"{sum_data.get('total_items', 0)} Items", "", "", "", f"Rs. {sum_data.get('total_stock_value', 0):,.2f}", ""])

        col_w = [width * 0.20, width * 0.10, width * 0.15, width * 0.11, width * 0.08, width * 0.12, width * 0.13, width * 0.11]
        t = Table(table_data, colWidths=col_w, repeatRows=1)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(f"#{cls.EMERALD}")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (3, 0), (3, -1), 'RIGHT'),
            ('ALIGN', (5, 0), (6, -1), 'RIGHT'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor("#F9F9F6")]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
            ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
            ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor("#F0FDF4")),
            ('PADDING', (0, 0), (-1, -1), 4),
        ]))
        elements.append(t)

    @classmethod
    def _build_pdf_movements(cls, elements, data, width, styles):
        movs = data.get("movements", [])[:100]
        table_data = [["Date", "Time", "Item Name", "Movement Type", "Quantity Delta", "Balance After", "UOM", "Reference"]]
        for m in movs:
            table_data.append([
                m["date"], m["time"], m["item_name"], m["movement_type"],
                f"{m['quantity_delta']:+.2f}", f"{m['balance_after']:.2f}",
                m["uom"], m["reference_note"][:30]
            ])

        col_w = [width * 0.10, width * 0.08, width * 0.20, width * 0.14, width * 0.11, width * 0.11, width * 0.08, width * 0.18]
        t = Table(table_data, colWidths=col_w, repeatRows=1)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(f"#{cls.EMERALD}")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (4, 0), (5, -1), 'RIGHT'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9F9F6")]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
            ('PADDING', (0, 0), (-1, -1), 4),
        ]))
        elements.append(t)

    @classmethod
    def _build_pdf_purchases(cls, elements, data, width, styles):
        pos = data.get("purchases", [])
        table_data = [["PO #", "Invoice #", "Supplier", "Purchase Date", "Status", "Subtotal", "Total (PKR)"]]
        for p in pos:
            table_data.append([
                p["purchase_number"], p["invoice_number"], p["supplier_name"],
                p["purchase_date"], p["status"], f"Rs. {p['subtotal']:,.2f}", f"Rs. {p['total_amount']:,.2f}"
            ])

        sum_data = data.get("summary", {})
        table_data.append(["TOTALS", "", "", f"{sum_data.get('total_purchase_count', 0)} POs", "", "", f"Rs. {sum_data.get('total_spend', 0):,.2f}"])

        col_w = [width * 0.12, width * 0.14, width * 0.24, width * 0.12, width * 0.10, width * 0.13, width * 0.15]
        t = Table(table_data, colWidths=col_w, repeatRows=1)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(f"#{cls.EMERALD}")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (5, 0), (6, -1), 'RIGHT'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor("#F9F9F6")]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
            ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
            ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor("#F0FDF4")),
            ('PADDING', (0, 0), (-1, -1), 4),
        ]))
        elements.append(t)

    @classmethod
    def _build_pdf_reconciliation(cls, elements, data, width, styles):
        rec = data.get("reconciliation", {})
        rev_rec = rec.get("revenue_reconciliation", {})
        cogs_rec = rec.get("cogs_reconciliation", {})
        exp_rec = rec.get("expense_reconciliation", {})

        table_data = [
            ["Audit Reconciliation Category", "Amount (PKR)", "Status", "Audit Notes"],
            ["Orders Gross Sales Total", f"Rs. {float(rev_rec.get('orders_gross_revenue', 0)):,.2f}", "", "Sales ledger"],
            ["Confirmed Payments Total", f"Rs. {float(rev_rec.get('payments_total', 0)):,.2f}", "", "Payment transactions"],
            ["Revenue Difference / Variance", f"Rs. {float(rev_rec.get('difference', 0)):,.2f}", rev_rec.get('status'), "Reconciled"],
            ["Order-Level Snapshotted COGS", f"Rs. {float(cogs_rec.get('order_cogs_total', 0)):,.2f}", "", "Orders.total_cogs"],
            ["Item-Level Snapshotted COGS", f"Rs. {float(cogs_rec.get('item_cogs_total', 0)):,.2f}", "", "OrderItem cost snapshot"],
            ["COGS Difference / Variance", f"Rs. {float(cogs_rec.get('difference', 0)):,.2f}", cogs_rec.get('status'), "Reconciled"],
            ["Active Expense Records Sum", f"Rs. {float(exp_rec.get('expense_records_sum', 0)):,.2f}", "", "Active Expense table"],
            ["P&L Summary Total", f"Rs. {float(exp_rec.get('expense_summary_total', 0)):,.2f}", "", "P&L Statement"],
            ["Expense Difference / Variance", f"Rs. {float(exp_rec.get('difference', 0)):,.2f}", exp_rec.get('status'), "Reconciled"],
        ]

        col_w = [width * 0.45, width * 0.22, width * 0.13, width * 0.20]
        t = Table(table_data, colWidths=col_w)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(f"#{cls.EMERALD}")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9F9F6")]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
            ('PADDING', (0, 0), (-1, -1), 5),
            ('FONTNAME', (0, 3), (-1, 3), 'Helvetica-Bold'),
            ('FONTNAME', (0, 6), (-1, 6), 'Helvetica-Bold'),
            ('FONTNAME', (0, 9), (-1, 9), 'Helvetica-Bold'),
        ]))
        elements.append(t)
